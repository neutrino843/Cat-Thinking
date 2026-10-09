"""Generate the small, project-owned PDF and DOCX files used by document-import E2E tests."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "e2e" / "fixtures" / "documents"
FIXED_TIME = datetime(2026, 10, 9, 9, 0, tzinfo=timezone.utc)


def set_run_font(run, name: str = "Microsoft YaHei", size: int = 11, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), name)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def normalize_docx_zip(path: Path) -> None:
    """Repack a DOCX with stable entry order and timestamps."""
    with zipfile.ZipFile(path, "r") as source:
        members = [(item.filename, source.read(item.filename)) for item in source.infolist()]
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name, payload in sorted(members):
            info = zipfile.ZipInfo(name, date_time=(2026, 10, 9, 9, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, payload)
    path.write_bytes(buffer.getvalue())


def create_docx() -> Path:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    path = OUTPUT / "structured-course.docx"
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    properties = document.core_properties
    properties.title = "文档解析课程"
    properties.subject = "Cat-Thinking DOCX import fixture"
    properties.author = "Cat-Thinking"
    properties.created = FIXED_TIME
    properties.modified = FIXED_TIME

    heading = document.add_heading("文档解析课程", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(heading.runs[0], size=20, bold=True)

    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(intro.add_run("这份文档用于验证 DOCX 标题、段落、制表符和表格提取。"))

    chapter = document.add_heading("第一章 来源与引用", level=2)
    set_run_font(chapter.runs[0], size=15, bold=True)
    paragraph = document.add_paragraph()
    set_run_font(paragraph.add_run("来源锚点"), bold=True)
    paragraph.add_run().add_tab()
    set_run_font(paragraph.add_run("把导图节点连接到原文字符区间。"))

    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    rows = [
        ("术语", "说明"),
        ("SourceDocument", "保存规范化原文和提取器版本"),
        ("SourceAnchor", "保存节点引用和字符区间"),
    ]
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, size=10, bold=row_index == 0)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")

    chapter = document.add_heading("第二章 可编辑导图", level=2)
    set_run_font(chapter.runs[0], size=15, bold=True)
    paragraph = document.add_paragraph()
    set_run_font(paragraph.add_run("导入后的节点仍可重命名、移动、撤销和重新导出。"))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("Cat-Thinking Fixture"), size=9)

    document.save(path)
    normalize_docx_zip(path)
    return path


def draw_pdf_header(pdf: canvas.Canvas, title: str, subtitle: str) -> None:
    pdf.setFillColorRGB(0.1, 0.16, 0.25)
    pdf.setFont("Helvetica-Bold", 22)
    pdf.drawString(72, 720, title)
    pdf.setFillColorRGB(0.28, 0.34, 0.42)
    pdf.setFont("Helvetica", 11)
    pdf.drawString(72, 697, subtitle)
    pdf.setStrokeColorRGB(0.78, 0.82, 0.86)
    pdf.line(72, 684, 540, 684)


def create_structured_pdf() -> Path:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    path = OUTPUT / "structured-course.pdf"
    pdf = canvas.Canvas(str(path), pagesize=letter, pageCompression=1, invariant=1)
    pdf.setTitle("PDF Import Course")
    pdf.setAuthor("Cat-Thinking")
    pdf.setSubject("Text-layer PDF import fixture")

    draw_pdf_header(pdf, "PDF Import Course", "Chapter 1  Source Anchors")
    pdf.setFillColorRGB(0.08, 0.08, 0.08)
    pdf.setFont("Helvetica", 12)
    pdf.drawString(72, 645, "Source anchors connect editable map nodes to page 1.")
    pdf.drawString(72, 620, "The body number 42 must remain in extracted text.")
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(72, 570, "Coverage checklist")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(88, 545, "1. Preserve reading order")
    pdf.drawString(88, 523, "2. Keep body numbers")
    pdf.setFont("Helvetica", 9)
    pdf.drawCentredString(306, 36, "1")
    pdf.showPage()

    draw_pdf_header(pdf, "Chapter 2  Editable Maps", "Page-aware extraction")
    pdf.setFillColorRGB(0.08, 0.08, 0.08)
    pdf.setFont("Helvetica", 12)
    pdf.drawString(72, 645, "Imported content remains editable after local extraction.")
    pdf.drawString(72, 620, "This sentence must be anchored to page 2.")
    pdf.setStrokeColorRGB(0.72, 0.76, 0.8)
    pdf.rect(72, 510, 468, 72, stroke=1, fill=0)
    pdf.line(230, 510, 230, 582)
    pdf.line(72, 546, 540, 546)
    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(82, 558, "Artifact")
    pdf.drawString(240, 558, "Purpose")
    pdf.setFont("Helvetica", 10)
    pdf.drawString(82, 524, "Mind map")
    pdf.drawString(240, 524, "Editable knowledge structure")
    pdf.setFont("Helvetica", 9)
    pdf.drawCentredString(306, 36, "2")
    pdf.save()
    return path


def create_scanned_pdf() -> Path:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    path = OUTPUT / "scanned-no-text.pdf"
    image = Image.new("RGB", (1275, 1650), "white")
    draw = ImageDraw.Draw(image)
    # Pillow's bundled font avoids host-specific font files and keeps the binary reproducible.
    title_font = ImageFont.load_default(size=56)
    body_font = ImageFont.load_default(size=32)
    draw.text((120, 180), "SCANNED PDF FIXTURE", fill="black", font=title_font)
    draw.text((120, 300), "This page is an image and has no PDF text layer.", fill="black", font=body_font)
    draw.rectangle((120, 410, 1155, 620), outline=(80, 100, 120), width=4)
    draw.text((165, 485), "Expected result: OCR-required error", fill=(30, 55, 75), font=body_font)

    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    buffer.seek(0)
    pdf = canvas.Canvas(str(path), pagesize=letter, pageCompression=1, invariant=1)
    pdf.setTitle("Scanned PDF Fixture")
    pdf.setAuthor("Cat-Thinking")
    pdf.drawImage(ImageReader(buffer), 0, 0, width=612, height=792, preserveAspectRatio=True)
    pdf.save()
    return path


def write_metadata() -> None:
    files = {
        "structured-course.pdf": {
            "mime": "application/pdf",
            "description": "Two-page PDF with a real text layer and edge page numbers",
        },
        "scanned-no-text.pdf": {
            "mime": "application/pdf",
            "description": "One-page image-only PDF without a text layer",
        },
        "structured-course.docx": {
            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "description": "OOXML document with headings, paragraphs, a tab, and a table",
        },
    }
    manifest = {"schemaVersion": 1, "generatedAt": "2026-10-09T09:00:00Z", "files": {}}
    for name, metadata in files.items():
        payload = (OUTPUT / name).read_bytes()
        manifest["files"][name] = {
            **metadata,
            "bytes": len(payload),
            "sha256": hashlib.sha256(payload).hexdigest(),
        }
    (OUTPUT / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    expected = {
        "schemaVersion": 1,
        "structuredPdf": {
            "file": "structured-course.pdf",
            "title": "PDF Import Course",
            "requiredText": [
                "The body number 42 must remain in extracted text.",
                "This sentence must be anchored to page 2.",
            ],
            "anchorPages": [1, 2],
        },
        "scannedPdf": {
            "file": "scanned-no-text.pdf",
            "errorContains": "尚未内建 OCR",
        },
        "structuredDocx": {
            "file": "structured-course.docx",
            "title": "文档解析课程",
            "requiredText": [
                "第一章 来源与引用",
                "把导图节点连接到原文字符区间。",
                "SourceDocument | 保存规范化原文和提取器版本",
                "第二章 可编辑导图",
            ],
        },
    }
    (OUTPUT / "expected.json").write_text(
        json.dumps(expected, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--format", choices=("docx", "pdf", "metadata", "all"), default="all")
    args = parser.parse_args()
    if args.format in ("docx", "all"):
        create_docx()
    if args.format in ("pdf", "all"):
        create_structured_pdf()
        create_scanned_pdf()
    if args.format in ("metadata", "all"):
        write_metadata()


if __name__ == "__main__":
    main()
